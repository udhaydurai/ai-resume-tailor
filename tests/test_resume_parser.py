"""Tests for resume parser module."""

import io

import pytest
from docx import Document

from src.resume_parser import parse_resume


class FakeUploadedFile:
    """Mimics a Streamlit UploadedFile."""

    def __init__(self, name: str, content: bytes):
        self.name = name
        self._buf = io.BytesIO(content)

    def read(self) -> bytes:
        return self._buf.read()

    def seek(self, pos: int) -> None:
        self._buf.seek(pos)


def test_parse_txt():
    text = "John Doe\nSoftware Engineer\n5 years experience"
    f = FakeUploadedFile("resume.txt", text.encode("utf-8"))
    result = parse_resume(f)
    assert "John Doe" in result
    assert "Software Engineer" in result


def test_parse_docx():
    doc = Document()
    doc.add_paragraph("Jane Smith")
    doc.add_paragraph("Product Manager")
    buf = io.BytesIO()
    doc.save(buf)
    f = FakeUploadedFile("resume.docx", buf.getvalue())
    result = parse_resume(f)
    assert "Jane Smith" in result
    assert "Product Manager" in result


def test_parse_unsupported():
    f = FakeUploadedFile("resume.html", b"<html></html>")
    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_resume(f)


def test_parse_pdf():
    """Smoke test: an empty-ish PDF should not crash."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    f = FakeUploadedFile("resume.pdf", buf.getvalue())
    result = parse_resume(f)
    assert isinstance(result, str)
