"""Generate downloadable cover letter documents (.docx)."""

from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def generate_cover_letter_docx(cover_letter_text: str, user_profile: dict) -> bytes:
    """Convert a plain-text cover letter into a .docx file.

    Returns the file content as bytes.
    """
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Contact header
    name = user_profile.get("name", "")
    email = user_profile.get("email", "")
    phone = user_profile.get("phone", "")
    linkedin = user_profile.get("linkedin", "")

    if name:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(name)
        run.bold = True
        run.font.size = Pt(14)
        run.font.name = "Calibri"

        contact_parts = [x for x in [email, phone, linkedin] if x]
        if contact_parts:
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run2 = p2.add_run(" | ".join(contact_parts))
            run2.font.size = Pt(10)
            run2.font.name = "Calibri"

        doc.add_paragraph("")  # spacer

    # Body paragraphs
    for line in cover_letter_text.strip().split("\n"):
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")
            continue
        p = doc.add_paragraph(stripped)
        for run in p.runs:
            run.font.size = Pt(11)
            run.font.name = "Calibri"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
