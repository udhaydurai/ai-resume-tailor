"""Generate downloadable resume documents (.docx and .pdf)."""

from __future__ import annotations

import io
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def generate_docx(resume_text: str) -> bytes:
    """Convert plain-text resume into a .docx file.

    Returns the file content as bytes.
    """
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    lines = resume_text.strip().split("\n")

    for line in lines:
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")
            continue

        # Detect section headers (ALL CAPS lines)
        if stripped.isupper() and len(stripped) > 2:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(12)
            run.font.name = "Calibri"
        # First line = name / contact header
        elif line == lines[0]:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(14)
            run.font.name = "Calibri"
        # Bullet points
        elif stripped.startswith(("- ", "• ", "* ")):
            p = doc.add_paragraph(stripped[2:], style="List Bullet")
            for run in p.runs:
                run.font.size = Pt(10.5)
                run.font.name = "Calibri"
        else:
            p = doc.add_paragraph(stripped)
            for run in p.runs:
                run.font.size = Pt(10.5)
                run.font.name = "Calibri"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_pdf(resume_text: str) -> bytes:
    """Convert plain-text resume into a PDF file.

    Returns the file content as bytes.
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
        leftMargin=0.8 * inch,
        rightMargin=0.8 * inch,
    )

    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            "ResumeHeader",
            parent=styles["Title"],
            fontSize=14,
            alignment=1,  # center
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            "SectionHeader",
            parent=styles["Heading2"],
            fontSize=12,
            spaceBefore=10,
            spaceAfter=4,
            textColor="black",
        )
    )
    styles.add(
        ParagraphStyle(
            "ResumeBody",
            parent=styles["Normal"],
            fontSize=10.5,
            leading=13,
            spaceAfter=2,
        )
    )
    styles.add(
        ParagraphStyle(
            "ResumeBullet",
            parent=styles["Normal"],
            fontSize=10.5,
            leading=13,
            leftIndent=18,
            spaceAfter=2,
            bulletIndent=6,
        )
    )

    story: list[Any] = []
    lines = resume_text.strip().split("\n")

    for i, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 6))
            continue

        if stripped.isupper() and len(stripped) > 2:
            story.append(Paragraph(escape(stripped), styles["SectionHeader"]))
        elif i == 0:
            story.append(Paragraph(escape(stripped), styles["ResumeHeader"]))
        elif stripped.startswith(("- ", "• ", "* ")):
            text = f"&bull; {escape(stripped[2:])}"
            story.append(Paragraph(text, styles["ResumeBullet"]))
        else:
            story.append(Paragraph(escape(stripped), styles["ResumeBody"]))

    doc.build(story)
    return buf.getvalue()
